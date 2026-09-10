"""Reusable editable 16:9 renderer. Scientific copy stays in deck.json."""
import argparse,math,subprocess,os
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
import pymupdf as fitz
from project import load,save,sha,digest,local,pages,utc

def inputs(root):
    deck=load(root/'deck.json');theme=load(root/'theme.json');assets=load(root/'assets.json',{})
    if not deck or not deck.get('slides'):raise ValueError('Author nonempty deck.json before building')
    ids=[s['id'] for s in deck['slides']]
    if len(ids)!=len(set(ids)):raise ValueError('Slide IDs must be unique')
    for k,a in assets.items():
        if sha(local(root,a['path']))!=a['sha256']:raise ValueError('Asset hash mismatch: '+k)
    return deck,theme,assets
def used(s):return [f['asset'] if isinstance(f,dict) else f for f in s.get('figures',[])]+([s['image']] if s.get('image') else [])
def fingerprints(deck,theme,assets):
    engine=sha(Path(__file__))
    return {s['id']:digest({'slide':s,'theme':theme,'engine_sha256':engine,'assets':{k:assets[k]['sha256'] for k in used(s)}}) for s in deck['slides']}
def changes(root):
    d,t,a=inputs(root);new=fingerprints(d,t,a);base=load(root/'build/baseline.json',{});old=base.get('slide_hashes',{});order=base.get('order',[])
    changed=[i for i,s in enumerate(d['slides'],1) if old.get(s['id'])!=new[s['id']] or i>len(order) or order[i-1]!=s['id']]
    return {'pages':changed,'slide_ids':[d['slides'][i-1]['id'] for i in changed],'removed_ids':[k for k in old if k not in new]}
def auto_boxes(keys,assets,root,area,groups=None):
    x,y,w,h=area;gap=.23
    if not keys:return []
    if groups is None:
        count=1 if w>10 or len(keys)<=2 else 2
        if len(keys)>7:count=2
        step=math.ceil(len(keys)/count);groups=[keys[i:i+step] for i in range(0,len(keys),step)]
    flat=[k for g in groups for k in g]
    if sorted(flat)!=sorted(keys):raise ValueError('groups must contain each figure once')
    rh=(h-gap*(len(groups)-1))/len(groups);result=[]
    for j,group in enumerate(groups):
        ratios=[]
        for k in group:
            iw,ih=Image.open(local(root,assets[k]['path'])).size;ratios.append(iw/ih)
        height=min(rh,(w-gap*(len(group)-1))/sum(ratios))
        if height<=0:raise ValueError('Too many panels for the automatic layout; supply figures[].box')
        widths=[r*height for r in ratios];cx=x+(w-sum(widths)-gap*(len(group)-1))/2
        for k,pw in zip(group,widths):
            result.append((k,[cx,y+j*(rh+gap)+(rh-height)/2,pw,height]));cx+=pw+gap
    return result
def build(root,selection=None):
    deck,theme,assets=inputs(root);all_slides=deck['slides']
    chosen=changes(root)['pages'] if selection=='changed' else pages(selection,len(all_slides))
    if not chosen:print('No changed slides; existing output retained');return
    prs=Presentation();prs.slide_width=Inches(16);prs.slide_height=Inches(9)
    records=[]
    def rect(sl,x,y,w,h,color):
        sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background()
        for e in sh._element.xpath('.//a:effectRef'):e.set('idx','0')
        return sh
    def text(sl,x,y,w,h,value,size,color,bold=False):
        sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        for j,part in enumerate(value.split('\n')):
            p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=part;p.font.name=theme['font'];p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.line_spacing=1.09;p.space_after=Pt(0);p.space_before=Pt(0)
            for r in p.runs:
                ea=OxmlElement('a:ea');ea.set('typeface',theme['font']);r._r.get_or_add_rPr().append(ea)
        return sh
    def line(sl,x1,y1,x2,y2,color,width=.8):
        sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
        for e in sh._element.xpath('.//a:effectRef'):e.set('idx','0')
    def image(sl,key,box,sid):
        a=assets[key];path=local(root,a['path']);iw,ih=Image.open(path).size;x,y,w,h=box;k=min(w/iw,h/ih);pw,ph=iw*k,ih*k;xx=x+(w-pw)/2;yy=y+(h-ph)/2
        sh=sl.shapes.add_picture(str(path),Inches(xx),Inches(yy),width=Inches(pw),height=Inches(ph));sh.name='Asset '+key
        sh._element.xpath('./p:nvPicPr/p:cNvPr')[0].set('descr',str({k:a[k] for k in ['kind','page','panel','readout','sha256'] if k in a}))
        records.append({'slide_id':sid,'asset':key,'sha256':a['sha256'],'kind':a['kind'],'box':[xx,yy,pw,ph]})
    navy=theme['navy'];burg=theme['burgundy'];gray=theme['gray']
    for number in chosen:
        s=all_slides[number-1];sl=prs.slides.add_slide(prs.slide_layouts[6]);kind=s.get('kind','result');sid=s['id']
        if kind=='locked-image':image(sl,s['image'],[0,0,16,9],sid)
        elif kind=='paper-cover':
            image(sl,s['image'],[.65,.55,14.7,7.9],sid)
        elif kind=='cover':
            if s.get('image'):image(sl,s['image'],[8.7,0,7.3,9],sid)
            text(sl,.8,2.0,7.4,2.6,s['title'],theme['title_pt'],navy,True)
            text(sl,.8,5.0,7.4,1.3,s.get('subtitle',''),22,gray)
            text(sl,.8,7.3,7.4,1.0,s.get('citation',''),18,gray)
        else:
            text(sl,.60,.22,14.8,.64,s['title'],theme['title_pt'],navy,True)
            text(sl,.61,.94,14.78,.41,s.get('condition',''),theme['condition_pt'],gray)
            line(sl,.60,1.45,15.4,1.45,theme['sage'],1.1);line(sl,.60,1.45,1.95,1.45,burg,1.8)
            layout=s.get('layout','side');wide=layout=='wide'
            area=[.65,1.72,14.7,3.5] if wide else [.65,1.77,8.05,5.36]
            figs=s.get('figures',[])
            if s.get('image'):figs=[s['image']]
            keys=[f['asset'] if isinstance(f,dict) else f for f in figs]
            if figs and all(isinstance(f,dict) and 'box' in f for f in figs):placements=[(f['asset'],f['box']) for f in figs]
            else:placements=auto_boxes(keys,assets,root,area,s.get('groups'))
            if kind=='result' and (not placements or any(assets[k]['kind']!='evidence' for k,b in placements)):raise ValueError('Result pages require original evidence assets: '+sid)
            for key,box in placements:image(sl,key,box,sid)
            nodes=s.get('chain',[])
            if kind=='concept' and not placements and not nodes:raise ValueError('Concept page needs an image or native chain: '+sid)
            if nodes:
                step=14.7/len(nodes)
                for j,node in enumerate(nodes):
                    x=.65+j*step;rect(sl,x,2.25,step-.28,1.85,theme['pale']);text(sl,x+.10,2.51,step-.48,.46,node['label'],20,navy,True);text(sl,x+.10,3.13,step-.48,.60,node.get('detail',''),18,gray)
                    if j+1<len(nodes):text(sl,x+step-.25,2.77,.24,.4,'›',22,burg,True)
            blocks=s.get('blocks',[])
            if len(blocks) not in [2,3]:raise ValueError('Use2 or3 explanation blocks: '+sid)
            for j,b in enumerate(blocks):
                if wide:
                    step=15.12/len(blocks);x=.65+j*step;y=5.35;w=step-.49;bh=1.28
                    if j:line(sl,x-.22,5.35,x-.22,7.12,'DEE6E3',.7)
                else:
                    step=5.40/len(blocks);x=9.13;y=1.77+j*step;w=6.18;bh=step-.6
                    if j:line(sl,x,y-.13,15.31,y-.13,'DEE6E3',.7)
                text(sl,x,y,w,.46,b['head'],theme['module_title_pt'],burg,True);text(sl,x,y+.53,w,bh,b['body'],theme['body_pt'],navy)
            rect(sl,.72,7.43,14.56,1.25,theme['pale']);rect(sl,.72,7.43,.055,1.25,burg)
            text(sl,.89,7.49,14.20,.47,'核心结论｜'+s['core'],theme['core_pt'],navy,True)
            text(sl,.89,7.99,14.20,.65,'补充说明｜'+s['supplement'],theme['supplement_pt'],gray)
        sl.notes_slide.notes_text_frame.text=s.get('notes','')
    out=root/('build/proof.pptx' if selection else 'build/deck.pptx');out.parent.mkdir(exist_ok=True);prs.save(out)
    meta={'created_at':utc(),'pptx_sha256':sha(out),'source_pages':chosen,'slide_ids':[all_slides[n-1]['id'] for n in chosen],'assets':records}
    save(out.with_suffix('.manifest.json'),meta)
    if not selection:save(root/'build/baseline.json',{'slide_hashes':fingerprints(deck,theme,assets),'order':[s['id'] for s in all_slides],'pptx_sha256':sha(out)})
    (out.parent/(out.stem+'.notes.md')).write_text('\n\n'.join(f"## {n}. {all_slides[n-1]['title']}\n\n{all_slides[n-1].get('notes','')}" for n in chosen),encoding='utf8')
    print(str(out))
def export(root,soffice,input_path=None):
    inp=input_path.resolve() if input_path else root/'build/deck.pptx'
    if not inp.exists():raise ValueError('PPTX missing')
    if not Path(soffice).is_file():raise ValueError('Renderer path does not exist; locate the installed soffice executable before retrying')
    out=root/'output';out.mkdir(exist_ok=True);profile=(root/'cache/lo-profile').resolve().as_uri()
    cmd=[str(soffice),'-env:UserInstallation='+profile,'--headless','--nologo','--nodefault','--nofirststartwizard','--norestore','--convert-to','pdf','--outdir',str(out),str(inp)]
    print('Exporting actual PPTX to PDF...',flush=True)
    r=subprocess.run(cmd,capture_output=True,text=True,encoding='utf8',errors='replace',timeout=240,creationflags=0x08000000 if os.name=='nt' else 0)
    target=out/(inp.stem+'.pdf')
    if r.returncode or not target.exists() or target.stat().st_mtime_ns<inp.stat().st_mtime_ns:raise RuntimeError(r.stdout+'\n'+r.stderr)
    save(root/f'build/{inp.stem}.export.json',{'pptx':str(inp),'pptx_sha256':sha(inp),'pdf':str(target),'pdf_sha256':sha(target),'created_at':utc()});print(str(target))
def raster(root,pdf,selection=None):
    doc=fitz.open(pdf);out=root/'output'/f'{pdf.stem}_renders';out.mkdir(parents=True,exist_ok=True);chosen=pages(selection,len(doc))
    for n in chosen:
        pg=doc[n-1];pg.get_pixmap(matrix=fitz.Matrix(1920/pg.rect.width,1920/pg.rect.width),alpha=False).save(out/f'slide-{n:03d}.png')
    save(out/'render_manifest.json',{'pdf_sha256':sha(pdf),'page_count':len(doc),'rendered_pages':chosen,'renders':{str(n):sha(out/f'slide-{n:03d}.png') for n in chosen}});print(str(out))
def main():
    ap=argparse.ArgumentParser();ap.add_argument('command',choices=['changes','build','export','raster']);ap.add_argument('--workspace',required=True,type=Path);ap.add_argument('--pages');ap.add_argument('--soffice',type=Path);ap.add_argument('--input',type=Path);ap.add_argument('--pdf',type=Path);a=ap.parse_args();root=a.workspace.resolve()
    if a.command=='changes':print(__import__('json').dumps(changes(root),ensure_ascii=False))
    elif a.command=='build':build(root,a.pages)
    elif a.command=='export':
        if not a.soffice:ap.error('--soffice required')
        export(root,a.soffice,a.input)
    else:raster(root,a.pdf or root/'output/deck.pdf',a.pages)
if __name__=='__main__':main()
