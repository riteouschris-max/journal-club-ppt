"""Check source/asset invariants and bind human visual review to current renders."""
import argparse,json,math,re,unicodedata
from pathlib import Path
from PIL import ImageFont
from pptx import Presentation
import pymupdf as fitz
from project import load,save,sha,digest,local,pages,utc,source
from render_deck import inputs,used,fingerprints

def story(s,mode):
    result={k:s.get(k) for k in ['id','kind','title','condition','core']};result['assets']=used(s)
    if mode=='typography':result.update(blocks=s.get('blocks'),supplement=s.get('supplement'),notes=s.get('notes'))
    if s.get('locked') or s.get('kind')=='locked-image':result={'whole_locked_slide':s}
    return result
def overlap(a,b):return min(a[0]+a[2],b[0]+b[2])-max(a[0],b[0])>.01 and min(a[1]+a[3],b[1]+b[3])-max(a[1],b[1])>.01
def audit(root,snapshot=False,mode='layout',replace=False,reviewed=None):
    d,t,a=inputs(root);slides=d['slides'];source(root);lock=root/'locks.json'
    if snapshot:
        if lock.exists() and not replace:raise ValueError('Snapshot exists; preserve it unless user authorized changing the locked content')
        save(lock,{'mode':mode,'slides':[story(s,mode) for s in slides],'asset_hashes':{k:a[k]['sha256'] for s in slides for k in used(s)}});print('Source invariants saved');return 0
    errors=[];warnings=[];old=load(lock)
    if old:
        if old['slides']!=[story(s,old['mode']) for s in slides]:errors.append('Locked slide order, claims, assets, or preserved slide changed')
        for k,h in old['asset_hashes'].items():
            if k not in a or a[k]['sha256']!=h:errors.append('Locked asset changed: '+k)
    inventory=load(root/'figure_inventory.json',[]);used_by={s['id']:set(used(s)) for s in slides}
    if any(s.get('kind','result')=='result' for s in slides) and not inventory:errors.append('Figure inventory missing')
    for panel in inventory:
        key=panel['id'];disposition=panel.get('disposition');keys=panel.get('assets',[key])
        if not panel.get('verified'):errors.append('Panel has not been source-checked: '+key)
        if panel.get('critical',True) and disposition!='slide':errors.append('Critical panel omitted: '+key)
        if disposition=='slide':
            targets=panel.get('slide_ids',[])
            if not targets or any(s not in used_by for s in targets):errors.append('Invalid panel destination: '+key)
            elif not set(keys).issubset(set().union(*(used_by[s] for s in targets))):errors.append('Panel asset missing from assigned pages: '+key)
        elif disposition not in ['notes','omitted_with_reason'] or not panel.get('reason'):errors.append('Unexplained panel omission: '+key)
    indexed=set(k for p in inventory for k in p.get('assets',[p['id']]))
    for s in slides:
        if s.get('kind','result')=='result':
            for k in used(s):
                if k not in indexed:errors.append('Evidence absent from panel inventory: '+k)
                src=local(root,a[k].get('source_path','source/paper.pdf'))
                if not src.exists() or a[k].get('source_sha256')!=sha(src):errors.append('Evidence source PDF hash not verified: '+k)
    ppt=root/'build/deck.pptx';prs=Presentation(ppt);manifest=load(root/'build/deck.manifest.json',{});baseline=load(root/'build/baseline.json',{})
    if manifest.get('pptx_sha256')!=sha(ppt):errors.append('PPTX does not match its build manifest')
    if baseline.get('slide_hashes')!=fingerprints(d,t,a):errors.append('Source changed after full build')
    if len(prs.slides)!=len(slides):errors.append('Slide count mismatch')
    norm=lambda s:re.sub(r'\s+','',unicodedata.normalize('NFKC',s))
    pdf=root/'output/deck.pdf';export=load(root/'build/deck.export.json',{});doc=None
    if pdf.exists():
        doc=fitz.open(pdf)
        if export.get('pptx_sha256')!=sha(ppt) or export.get('pdf_sha256')!=sha(pdf):errors.append('PDF export is stale or unbound')
        if len(doc)!=len(slides):errors.append('PDF page count mismatch')
    else:errors.append('Final PDF missing')
    fp=Path(t.get('font_file','C:/Windows/Fonts/msyh.ttc'));font_cache={}
    if not fp.exists():warnings.append('Font metrics unavailable: review text wrapping visually or supply theme.font_file')
    for i,sl in enumerate(prs.slides):
        if i>=len(slides):break
        s=slides[i];kind=s.get('kind','result');text_boxes=[];image_boxes=[];actual=[]
        native=[]
        for sh in sl.shapes:
            b=[sh.left/914400,sh.top/914400,sh.width/914400,sh.height/914400]
            if min(b)<0 or b[0]+b[2]>16.001 or b[1]+b[3]>9.001:errors.append(f'Page{i+1}: object outside slide')
            if hasattr(sh,'image'):
                key=sh.name.removeprefix('Asset ');actual.append(key)
                if key not in a or __import__('hashlib').sha256(sh.image.blob).hexdigest()!=a[key]['sha256']:errors.append(f'Page{i+1}: embedded asset mismatch')
                if kind=='result':
                    image_boxes.append(b)
                    if any(getattr(sh,'crop_'+v)!=0 for v in ['left','right','top','bottom']):errors.append(f'Page{i+1}: evidence cropped in PPT')
            if sh.has_text_frame and sh.text:
                native.append(sh.text);text_boxes.append(b)
                estimate=0
                for p in sh.text_frame.paragraphs:
                    size=p.font.size.pt if p.font.size else None
                    if size and size<18:errors.append(f'Page{i+1}: authored text below18pt')
                    if size and fp.exists():
                        if size not in font_cache:font_cache[size]=ImageFont.truetype(str(fp),round(size*4))
                        width=font_cache[size].getlength(p.text)/4/72;rows=max(1,math.ceil(width/max(.01,b[2])))
                        estimate+=rows*size/72*1.16
                if estimate>b[3]+.10:errors.append(f'Page{i+1}: estimated text overflow: '+sh.text[:36])
        if sorted(actual)!=sorted(used(s)):errors.append(f'Page{i+1}: asset assignment differs from source')
        if kind=='result':
            if any(overlap(x,y) for j,x in enumerate(image_boxes) for y in image_boxes[j+1:]):errors.append(f'Page{i+1}: evidence images overlap')
            if any(overlap(x,y) for x in image_boxes for y in text_boxes):errors.append(f'Page{i+1}: evidence/text overlap')
        if any(overlap(x,y) for j,x in enumerate(text_boxes) for y in text_boxes[j+1:]):errors.append(f'Page{i+1}: text boxes overlap')
        if doc and i<len(doc) and any(norm(x) not in norm(doc[i].get_text()) for x in native):errors.append(f'Page{i+1}: native text missing from PDF')
        if not s.get('notes','').strip():warnings.append(f'Page{i+1}: no speaker notes')
    rd=root/'output/deck_renders';rm=load(rd/'render_manifest.json',{});previous=load(root/'qa/visual_review.json',{}).get('pages',{})
    accepted={};current=rm.get('renders',{})
    render_ok=doc is not None and rm.get('pdf_sha256')==sha(pdf)
    if reviewed and not render_ok:errors.append('Cannot record review against stale/missing renders')
    reviewed_pages=pages(reviewed,len(slides)) if reviewed else []
    if render_ok:
        for n in range(1,len(slides)+1):
            p=rd/f'slide-{n:03d}.png';h=current.get(str(n))
            if h and p.exists() and sha(p)==h:
                if n in reviewed_pages:accepted[str(n)]={'png_sha256':h,'reviewed_at':utc()}
                elif previous.get(str(n),{}).get('png_sha256')==h:accepted[str(n)]=previous[str(n)]
    visual_ok=len(accepted)==len(slides)
    if reviewed and not errors:save(root/'qa/visual_review.json',{'pptx_sha256':sha(ppt),'pdf_sha256':sha(pdf),'pages':accepted})
    result={'status':'PASS' if not errors and visual_ok else 'FAIL' if errors else 'NEEDS_VISUAL_REVIEW','errors':errors,'warnings':warnings,'slides':len(slides),'visual_reviewed_pages':[int(k) for k in accepted],'pptx_sha256':sha(ppt),'pdf_sha256':sha(pdf) if pdf.exists() else None}
    save(root/'qa/report.json',result);print(json.dumps(result,ensure_ascii=False,indent=2));return 1 if errors else 0
def main():
    ap=argparse.ArgumentParser();ap.add_argument('--workspace',required=True,type=Path);ap.add_argument('--snapshot',action='store_true');ap.add_argument('--mode',choices=['layout','typography'],default='layout');ap.add_argument('--replace-snapshot',action='store_true');ap.add_argument('--reviewed-pages');a=ap.parse_args()
    raise SystemExit(audit(a.workspace.resolve(),a.snapshot,a.mode,a.replace_snapshot,a.reviewed_pages))
if __name__=='__main__':main()
